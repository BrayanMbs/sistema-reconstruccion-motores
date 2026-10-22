"""Regenera los capítulos derivados y el Word desde fuentes Markdown locales.

Uso desde raíz: python docs/manual-tecnico/herramientas/generar_manual.py
Dependencias: python-docx, Pillow. No accede a bases de datos ni a la red.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'docs/manual-tecnico'
def read(path): return (ROOT / path).read_text(encoding='utf-8-sig')
def write(name, content):
    p = OUT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content.strip() + '\n', encoding='utf-8')
def link(path): return f'[{path}](../../{path})'
def table(headers, rows):
    esc = lambda x: str(x).replace('|', '\\|').replace('\n', ' ')
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] + ['| ' + ' | '.join(map(esc, row)) + ' |' for row in rows])

def inventories():
    # Extrae cadenas de rutas y métodos; no evalúa JavaScript.
    mounts = {'admin':'/api/admin','administrative':'/api/administrativo','operational':'/api/operativo','auth':'/api/auth','cashier':'/api/cashier','inventory':'/api/inventory','part-request':'/api/part-requests','public':'/api/public','health':'/api'}
    roles = {'admin':'ADMIN','administrative':'ADMINISTRATIVE','operational':'OPERATOR','auth':'Mixto: login público; me/cambio autenticados','cashier':'ADMIN, CASHIER','inventory':'ADMIN, INVENTORY','part-request':'OPERATOR, ADMIN, INVENTORY; restricciones por acción','public':'Público','health':'Público'}
    rows=[]
    for key, prefix in mounts.items():
        path=f'backend/src/routes/{key}.routes.ts'; src=read(path)
        for m in re.finditer(r'\w+Router\.(get|post|put|patch|delete|route)\("([^"]+)"([^;]*)(?:;|$)',src):
            method, route, tail=m.groups()
            methods=re.findall(r'\.(get|post|put|patch|delete)\(',tail) if method=='route' else [method]
            line=src.count('\n',0,m.start())+1
            for verb in methods: rows.append((verb.upper(),prefix+route,roles[key],f'{path}:{line}'))
    rows.insert(0,('GET','/health','Público','backend/src/app.ts'))
    write('anexo-endpoints.md', '# Catálogo de endpoints\n\nInventario estático de las declaraciones de rutas del commit auditado. Los permisos particulares y contratos se explican en [API](10-api-rest.md). `ADMIN` no hereda acceso a las áreas administrativa y operativa.\n\n'+table(['Método','Ruta','Acceso del router','Evidencia'],rows))
    schema=[]; indexes=[]; relations=[]; entities=[]
    for p in sorted((ROOT/'backend/database/migrations').glob('*.sql')):
        if p.name.startswith('004'): continue  # No copiar datos personales del seed.
        src=p.read_text(encoding='utf-8-sig')
        for m in re.finditer(r'CREATE TABLE IF NOT EXISTS (?:public\.)?(\w+)\s*\((.*?)\n\);',src,re.S):
            name,body=m.groups(); entities.append(name); schema.append(f'## {name}\n\nFuente: {link(p.relative_to(ROOT).as_posix())}. Definición inicial, complementada por las alteraciones listadas abajo.\n\n```sql\nCREATE TABLE {name} (\n{body}\n);\n```')
            for fk in re.finditer(r'(\w+)\s+uuid([^,\n]*?)REFERENCES (\w+)\(id\)',body): relations.append((fk[3],name,fk[1], '||' if 'NOT NULL' in fk[2] else '|o'))
        for m in re.finditer(r'ALTER TABLE (?:public\.)?(\w+)(.*?);',src,re.S):
            schema.append(f'### Modificación posterior de {m[1]}\n\nFuente: `{p.relative_to(ROOT).as_posix()}`.\n\n```sql\n{m[0]}\n```')
            for fk in re.finditer(r'(\w+)\s+uuid([^,\n]*?)REFERENCES (\w+)\(id\)',m[2]): relations.append((fk[3],m[1],fk[1], '||' if 'NOT NULL' in fk[2] else '|o'))
        for m in re.finditer(r'CREATE (?:UNIQUE )?INDEX IF NOT EXISTS .*?;',src,re.S): indexes.append(f'-- {p.name}\n{m[0]}')
    write('anexo-esquema.md','# Diccionario SQL y restricciones\n\nEsquema declarado, no inspección de una base desplegada. **NO VERIFICADO**: esquema efectivo, permisos y contenido de cualquier instancia. Las alteraciones posteriores prevalecen sobre las definiciones iniciales. No se reproduce el seed con datos personales.\n\n'+'\n\n'.join(schema)+'\n\n## Índices declarados\n\n```sql\n'+'\n'.join(indexes)+'\n```')
    entity_blocks=[]
    for name in entities:
        key='string setting_key PK' if name=='app_settings' else ('uuid work_order_id PK,FK\n        uuid inventory_item_id PK,FK' if name=='work_order_inventory' else ('bigint id PK' if name=='audit_events' else 'uuid id PK'))
        entity_blocks.append(f'    {name} {{\n        {key}\n    }}')
    er='erDiagram\n'+ '\n'.join(entity_blocks)+'\n'+ '\n'.join(f'    {a} {card}--o{{ {b} : {col}' for a,b,col,card in dict.fromkeys(relations))
    write('diagramas/04-entidad-relacion.mmd',er)
    pages=[p.relative_to(ROOT).as_posix() for p in (ROOT/'frontend/src/app').rglob('page.tsx')]
    write('anexo-rutas-frontend.md','# Rutas del frontend\n\nLos grupos entre paréntesis no forman parte de la URL. Los segmentos `[id]` son dinámicos. La protección efectiva de datos se impone en la API.\n\n'+table(['URL','Archivo'],[('/'+ '/'.join(s for s in p.split('/')[3:-1] if not s.startswith('(')),p) for p in sorted(pages)]))
    tests=[]
    for base in ['backend/tests','frontend/src']:
        for p in sorted((ROOT/base).rglob('*test.*')):
            if p.suffix in ['.ts','.tsx']: tests.append((p.relative_to(ROOT).as_posix(), 'Backend' if base.startswith('backend') else 'Frontend'))
    write('anexo-pruebas.md','# Inventario de pruebas\n\nLos nombres indican intención de cobertura, no cobertura porcentual ni validación de infraestructura real. Resultados de ejecución: [capítulo de calidad](12-pruebas-calidad.md).\n\n'+table(['Archivo','Área'],tests))

def diagram_images():
    expected=['01-arquitectura-general','02-backend','03-frontend','04-entidad-relacion','04a-ordenes-personal','04b-inventario-solicitudes','04c-finanzas-seguridad','05-autenticacion','06-orden-trabajo','07-git-ci','08-despliegue']
    if all((OUT/f'diagramas/{name}.png').exists() for name in expected):return
    raise RuntimeError('Renderice primero las fuentes Mermaid con herramientas/renderizar_diagramas.mjs')

def inline(p,text):
    # Conserva destinos de referencias en texto para impresión.
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
    for n,part in enumerate(re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text)):
        r=p.add_run(part.strip('`*') if part.startswith(('`','**')) else part)
        if part.startswith('`'): r.font.name='Consolas';r.font.size=Pt(9)
        if part.startswith('**'): r.bold=True

def export_word():
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
    for name in ['Heading 1','Heading 2','Heading 3','Title']:
        doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor.from_string('000000')
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    sec.header.paragraphs[0].text='Sistema de Reconstrucción de Motores | Manual técnico interno'
    p=sec.footer.paragraphs[0];p.alignment=2;p.add_run('Versión 1.0 | Página ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.add_paragraph('MANUAL TÉCNICO INTERNO PARA DESARROLLADORES','Title')
    doc.add_paragraph('Sistema de Reconstrucción de Motores','Subtitle')
    for t in ['Curso: Análisis de Sistemas II','Tipo de documento: Documentación técnica del software','Versión del documento: 1.0','Fecha de elaboración: 22 de octubre de 2026','Universidad: [Completar]','Integrantes: [Completar]','Catedrático: [Completar]','Sección: [Completar]']: doc.add_paragraph(t)
    doc.add_heading('Control de versiones del documento',1)
    t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,['Versión','Fecha','Referencia']):c.text=s
    for c,s in zip(t.add_row().cells,['1.0','22/10/2026','develop · 9d158a875bacc5b9946501fde7466e47105ee9d6']): c.text=s
    doc.add_paragraph('Auditoría estática del código y verificaciones locales. Los resultados y límites de comprobación constan en el capítulo de calidad y en la auditoría técnica.')
    doc.add_heading('Índice',1);p=doc.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');p._p.append(f)
    doc.add_paragraph('En Word, seleccione el índice y elija Actualizar toda la tabla. Los títulos también están disponibles en el panel de navegación.')
    files=sorted(OUT.glob('[0-1][0-9]-*.md'))+[OUT/'auditoria-tecnica.md']+sorted(OUT.glob('anexo-*.md'))
    for file in files:
        lines=file.read_text(encoding='utf-8').splitlines();i=0;code=False
        while i<len(lines):
            line=lines[i];i+=1
            if line.startswith('```'):code=not code;continue
            if code:
                # Wrap long SQL/TS lines without deleting content.
                import textwrap
                for chunk in textwrap.wrap(line,width=92,replace_whitespace=False,drop_whitespace=False) or ['']:
                    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    r=p.add_run(chunk);r.font.name='Consolas';r.font.size=Pt(8)
                continue
            if line.startswith('|'):
                block=[line]
                while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
                data=[[v.strip().replace('\\|','/') for v in s.strip('|').split('|')] for s in block if not re.match(r'^\|[ :|\-]+$',s)]
                if not data:continue
                cols=len(data[0]);t=doc.add_table(rows=0,cols=cols);t.style='Table Grid';t.autofit=False
                for vals in data:
                    row=t.add_row()
                    for c,v in zip(row.cells,vals):
                        c.width=Inches(6.5/cols);inline(c.paragraphs[0],v)
                        for r in c.paragraphs[0].runs:r.font.size=Pt(8)
                    prop=row._tr.get_or_add_trPr();pr=OxmlElement('w:cantSplit');prop.append(pr)
                prop=t.rows[0]._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
                for cell in t.rows[0].cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
                continue
            if line.startswith('#'):
                level=len(line)-len(line.lstrip('#'));doc.add_heading(line.lstrip('# ').replace('`',''),min(level,3));continue
            image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            if image:
                path=file.parent/image[2]
                if path.exists():
                    with Image.open(path) as im:w,h=im.size
                    width=min(6.1,6.7*w/h)
                    doc.add_picture(str(path),width=Inches(width))
                continue
            if not line.strip():continue
            p=doc.add_paragraph(style='List Bullet' if line.startswith('- ') else None);inline(p,line[2:] if line.startswith('- ') else line)
    settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='Manual técnico interno para desarrolladores del Sistema de Reconstrucción de Motores'
    target=OUT/'Manual_Tecnico_Sistema_Reconstruccion_Motores.docx';doc.save(target)
    print(f'Word generado: {target}; capítulos y anexos: {len(files)}')

if __name__=='__main__':
    inventories();diagram_images();export_word()
